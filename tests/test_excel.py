"""Tests for Excel export service."""
import io
import pytest
from openpyxl import load_workbook

from app.models.entities import InsuranceRecord, SyncRunSummary, ProcessingError
from app.services.excel_export_service import build_workbook


@pytest.fixture
def sample_records():
    return [
        InsuranceRecord(
            insurance_type="car",
            insurer_company="Migdal",
            policy_number="111222",
            insured_name="John Doe",
            premium_monthly=300.0,
            premium_yearly=3600.0,
            currency="ILS",
            start_date="2024-01-01",
            end_date="2024-12-31",
            needs_review=False,
            extraction_confidence=0.85,
            classification_confidence=0.9,
        ),
        InsuranceRecord(
            insurance_type="health",
            insurer_company="Harel",
            policy_number="333444",
            insured_name="Jane Doe",
            premium_monthly=150.0,
            premium_monthly_derived=True,
            premium_yearly=1800.0,
            currency="ILS",
            needs_review=True,
            extraction_confidence=0.35,
            classification_confidence=0.5,
            extraction_notes="Low confidence",
        ),
    ]


@pytest.fixture
def sample_runs():
    return [
        SyncRunSummary(
            status="completed",
            messages_found=20,
            messages_new=5,
            records_created=2,
        )
    ]


@pytest.fixture
def sample_errors():
    return [
        ProcessingError(
            stage="fetch",
            error_type="HttpError",
            error_msg="Rate limit exceeded",
        )
    ]


def _wb_bytes(records, runs, errors):
    wb = build_workbook(records, runs, errors)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_workbook_has_four_sheets(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    assert set(wb.sheetnames) == {"All Policies", "Monthly Payments", "Needs Review", "Sync Log"}


def test_all_policies_has_header_row(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["All Policies"]
    # Header row at row 1
    assert ws.cell(row=1, column=1).value == "Type"
    assert ws.cell(row=1, column=2).value == "Insurer"


def test_all_policies_data_rows(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["All Policies"]
    # 2 records → rows 2 and 3
    assert ws.cell(row=2, column=1).value == "car"
    assert ws.cell(row=2, column=2).value == "Migdal"
    assert ws.cell(row=3, column=1).value == "health"


def test_monthly_payments_only_includes_monthly_records(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["Monthly Payments"]
    # Both records have monthly premium
    data_rows = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(data_rows) >= 2


def test_needs_review_sheet_contains_only_review_records(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["Needs Review"]
    # Only 1 record needs review
    data_rows = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(data_rows) == 1
    assert data_rows[0] == "health"


def test_sync_log_sheet_has_run(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["Sync Log"]
    assert ws.cell(row=2, column=2).value == "completed"
    assert ws.cell(row=2, column=5).value == 20


def test_build_workbook_empty_records():
    """Should not crash with no data."""
    wb = build_workbook([], [], [])
    assert len(wb.sheetnames) == 4


def test_monthly_payments_total_row(sample_records, sample_runs, sample_errors):
    buf = _wb_bytes(sample_records, sample_runs, sample_errors)
    wb = load_workbook(buf)
    ws = wb["Monthly Payments"]
    # Find total row (labelled "TOTAL")
    total_row = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "TOTAL":
                total_row = cell.row
    assert total_row is not None
    total_val = ws.cell(row=total_row, column=5).value
    assert total_val == pytest.approx(450.0, rel=0.01)  # 300 + 150
