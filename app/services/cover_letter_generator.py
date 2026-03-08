"""Cover letter generation - only when explicitly required by the job."""

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.config import Config
from app.core.database import get_db
from app.models.job import ExtractedJob
from app.services.resume_generator import sanitize_filename, convert_to_pdf

logger = logging.getLogger(__name__)


def needs_cover_letter(job: ExtractedJob) -> bool:
    """Check if the job posting explicitly requests a cover letter."""
    text = f"{job.raw_message} {job.description} {job.requirements}".lower()
    cover_letter_signals = [
        "cover letter",
        "מכתב מלווה",
        "מכתב מקדים",
        "מכתב נלווה",
        "covering letter",
    ]
    return any(signal in text for signal in cover_letter_signals)


def build_cover_letter_docx(job: ExtractedJob, output_path: Path) -> Path:
    """Build a cover letter DOCX from source-of-truth data only."""
    config = Config.get()
    candidate = config.candidate

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Date
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    date_p.add_run(datetime.now().strftime("%B %d, %Y"))

    doc.add_paragraph()  # Spacing

    # Greeting
    if job.company:
        doc.add_paragraph(f"Dear {job.company} Hiring Team,")
    else:
        doc.add_paragraph("Dear Hiring Manager,")

    # Body
    title = job.job_title or "the open position"
    body = (
        f"I am writing to express my interest in the {title} position"
        f"{' at ' + job.company if job.company else ''}. "
        f"Please find my resume attached for your review.\n\n"
    )

    # Add summary from source of truth
    if candidate.get("summary"):
        body += candidate["summary"].strip() + "\n\n"

    body += (
        "I am eager to bring my skills and enthusiasm to your team "
        "and would welcome the opportunity to discuss how I can contribute.\n\n"
        "Thank you for your time and consideration."
    )

    doc.add_paragraph(body)

    # Closing
    doc.add_paragraph()
    doc.add_paragraph("Best regards,")
    closing = doc.add_paragraph()
    name_run = closing.add_run(candidate.get("full_name", ""))
    name_run.bold = True

    contact_lines = []
    if candidate.get("phone"):
        contact_lines.append(candidate["phone"])
    if candidate.get("email"):
        contact_lines.append(candidate["email"])
    if contact_lines:
        doc.add_paragraph("\n".join(contact_lines))

    doc.save(str(output_path))
    logger.info(f"Cover letter DOCX saved: {output_path}")
    return output_path


def generate_cover_letter(
    job: ExtractedJob, output_folder: Path, job_id: int | None = None
) -> dict | None:
    """Generate cover letter if needed. Returns paths dict or None."""
    if not needs_cover_letter(job):
        logger.info("Cover letter not required for this job")
        return None

    config = Config.get()
    candidate = config.candidate
    name = sanitize_filename(candidate.get("full_name", "CoverLetter"))
    company = sanitize_filename(job.company or "Company")
    date_str = datetime.now().strftime("%Y%m%d")

    docx_path = Path(output_folder) / f"Cover_Letter_{name}_{company}_{date_str}.docx"
    build_cover_letter_docx(job, docx_path)

    pdf_path = convert_to_pdf(docx_path)

    if job_id:
        with get_db() as db:
            db.execute(
                "UPDATE jobs SET generated_cover_letter_docx = ?, generated_cover_letter_pdf = ? WHERE id = ?",
                (str(docx_path), str(pdf_path) if pdf_path else "", job_id),
            )
            for fmt, path in [("docx", docx_path), ("pdf", pdf_path)]:
                if path:
                    db.execute(
                        "INSERT INTO generated_files (job_id, file_type, file_format, file_path) VALUES (?, ?, ?, ?)",
                        (job_id, "cover_letter", fmt, str(path)),
                    )

    return {
        "docx": str(docx_path),
        "pdf": str(pdf_path) if pdf_path else None,
    }
