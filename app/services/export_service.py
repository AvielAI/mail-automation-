"""Export service - Excel (.xlsx) and Word (.docx) reports."""

import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.config import GENERATED_DIR
from app.core.database import get_db

logger = logging.getLogger(__name__)

EXCEL_COLUMNS = [
    ("capture_date", "Capture Date"),
    ("sent_date", "Sent Date"),
    ("whatsapp_group", "WhatsApp Group"),
    ("job_title", "Job Title"),
    ("company", "Company"),
    ("location", "Location"),
    ("role_category", "Role Category"),
    ("selected_application_profile", "Application Profile"),
    ("apply_method", "Apply Method"),
    ("apply_email", "Apply Email"),
    ("apply_link", "Apply Link"),
    ("generated_resume_docx", "Resume DOCX"),
    ("generated_resume_pdf", "Resume PDF"),
    ("submission_status", "Status"),
    ("notes", "Notes"),
]


def export_to_excel(filters: dict | None = None) -> bytes:
    """Export jobs to Excel bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Job Applications"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A3C6E", end_color="1A3C6E", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Write headers
    for col_idx, (_, header) in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Fetch data
    jobs = _fetch_jobs(filters)

    # Write rows
    for row_idx, job in enumerate(jobs, 2):
        for col_idx, (field, _) in enumerate(EXCEL_COLUMNS, 1):
            value = job.get(field, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value or "")
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True)

    # Auto-width columns
    for col_idx, (_, header) in enumerate(EXCEL_COLUMNS, 1):
        max_len = len(header)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 50))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 2

    # Save to bytes
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_word(filters: dict | None = None) -> bytes:
    """Export jobs to Word report bytes."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading("Job Application Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font.size = Pt(9)

    doc.add_paragraph()

    # Fetch data
    jobs = _fetch_jobs(filters)

    if not jobs:
        doc.add_paragraph("No job applications found.")
    else:
        for i, job in enumerate(jobs, 1):
            doc.add_heading(f"{i}. {job.get('job_title', 'Unknown Position')}", level=2)

            details = [
                ("Company", job.get("company", "")),
                ("WhatsApp Group", job.get("whatsapp_group", "")),
                ("Capture Date", job.get("capture_date", "")),
                ("Location", job.get("location", "")),
                ("Role Category", job.get("role_category", "")),
                ("Apply Method", job.get("apply_method", "")),
                ("Status", job.get("submission_status", "")),
                ("Sent Date", job.get("sent_date", "")),
            ]

            for label, value in details:
                if value:
                    p = doc.add_paragraph()
                    p.add_run(f"{label}: ").bold = True
                    p.add_run(str(value))

            if job.get("notes"):
                p = doc.add_paragraph()
                p.add_run("Notes: ").bold = True
                p.add_run(job["notes"])

            if i < len(jobs):
                doc.add_page_break()

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def save_export_file(data: bytes, filename: str) -> Path:
    """Save export file to generated directory."""
    exports_dir = GENERATED_DIR / "exports"
    exports_dir.mkdir(exist_ok=True)
    path = exports_dir / filename
    path.write_bytes(data)
    logger.info(f"Export saved: {path}")
    return path


def _fetch_jobs(filters: dict | None = None) -> list[dict]:
    """Fetch jobs from database with optional filters."""
    query = "SELECT * FROM jobs ORDER BY capture_date DESC"
    params = []

    if filters:
        conditions = []
        if filters.get("status"):
            conditions.append("submission_status = ?")
            params.append(filters["status"])
        if filters.get("group"):
            conditions.append("whatsapp_group = ?")
            params.append(filters["group"])
        if filters.get("from_date"):
            conditions.append("capture_date >= ?")
            params.append(filters["from_date"])
        if filters.get("to_date"):
            conditions.append("capture_date <= ?")
            params.append(filters["to_date"])
        if conditions:
            query = f"SELECT * FROM jobs WHERE {' AND '.join(conditions)} ORDER BY capture_date DESC"

    with get_db() as db:
        rows = db.execute(query, params).fetchall()
        return [dict(row) for row in rows]
