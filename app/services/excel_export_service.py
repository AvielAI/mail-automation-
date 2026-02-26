"""Excel workbook builder — openpyxl, four sheets."""
from __future__ import annotations

import logging
from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter

from app.models.entities import InsuranceRecord, SyncRunSummary, ProcessingError

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Colours
# ---------------------------------------------------------------------------
HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
ALT_FILL      = PatternFill("solid", fgColor="DCE6F1")
REVIEW_FILL   = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL    = PatternFill("solid", fgColor="FFDCE1")
HEADER_FONT   = Font(color="FFFFFF", bold=True, size=10)
DATA_FONT     = Font(size=10)
THIN_BORDER   = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def build_workbook(
    records: list[InsuranceRecord],
    sync_runs: list[SyncRunSummary],
    errors: list[ProcessingError],
) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    _build_all_policies(wb, records)
    _build_monthly_payments(wb, records)
    _build_needs_review(wb, records)
    _build_sync_log(wb, sync_runs, errors)

    return wb


# ---------------------------------------------------------------------------
# Sheet: All Policies
# ---------------------------------------------------------------------------

_ALL_POLICIES_HEADERS = [
    ("Type", 14),
    ("Insurer", 22),
    ("Policy #", 16),
    ("Insured Name", 22),
    ("ID (masked)", 12),
    ("Start Date", 12),
    ("End Date", 12),
    ("Renewal Date", 12),
    ("Monthly (₪)", 13),
    ("Monthly Derived?", 16),
    ("Annual (₪)", 13),
    ("Annual Derived?", 15),
    ("Freq.", 10),
    ("Currency", 10),
    ("Amount Due", 13),
    ("Payment Status", 15),
    ("Vehicle #", 12),
    ("Property Address", 28),
    ("Agent", 18),
    ("Agency", 18),
    ("Email Subject", 35),
    ("Email Date", 14),
    ("Sender", 25),
    ("Attachment", 22),
    ("Extraction Conf.", 14),
    ("Class. Conf.", 12),
    ("Needs Review?", 14),
    ("Notes", 40),
    ("Record ID", 16),
]


def _build_all_policies(wb: Workbook, records: list[InsuranceRecord]) -> None:
    ws = wb.create_sheet("All Policies")
    headers = [h for h, _ in _ALL_POLICIES_HEADERS]
    widths = [w for _, w in _ALL_POLICIES_HEADERS]
    _write_headers(ws, headers, widths)

    for i, r in enumerate(records, start=2):
        row = [
            r.insurance_type,
            r.insurer_company,
            r.policy_number,
            r.insured_name,
            r.id_number_masked,
            r.start_date,
            r.end_date,
            r.renewal_date,
            r.premium_monthly,
            "Yes" if r.premium_monthly_derived else "No",
            r.premium_yearly,
            "Yes" if r.premium_yearly_derived else "No",
            r.payment_frequency,
            r.currency,
            r.amount_due,
            r.payment_status,
            r.vehicle_number,
            r.property_address,
            r.agent_name,
            r.agency_name,
            r.source_email_subject,
            r.source_email_date,
            r.source_email_sender,
            r.source_attachment_filename,
            r.extraction_confidence,
            r.classification_confidence,
            "Yes" if r.needs_review else "No",
            r.extraction_notes,
            r.record_id,
        ]
        _write_row(ws, i, row, alt=(i % 2 == 0),
                   highlight=REVIEW_FILL if r.needs_review else None)
        # Number formats
        for col_offset, col_name in [(9, "monthly"), (11, "annual"), (15, "due")]:
            cell = ws.cell(row=i, column=col_offset)
            if cell.value is not None:
                cell.number_format = '#,##0.00'
        for col_offset in [25, 26]:
            cell = ws.cell(row=i, column=col_offset)
            if cell.value is not None:
                cell.number_format = "0.00%"
                cell.value = cell.value  # already 0–1 float

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Sheet: Monthly Payments
# ---------------------------------------------------------------------------

_MONTHLY_HEADERS = [
    ("Type", 14),
    ("Insurer", 22),
    ("Policy #", 16),
    ("Insured Name", 22),
    ("Monthly (₪)", 14),
    ("Derived?", 10),
    ("Currency", 10),
    ("Renewal Date", 14),
    ("Needs Review?", 14),
]


def _build_monthly_payments(wb: Workbook, records: list[InsuranceRecord]) -> None:
    ws = wb.create_sheet("Monthly Payments")
    headers = [h for h, _ in _MONTHLY_HEADERS]
    widths = [w for _, w in _MONTHLY_HEADERS]
    _write_headers(ws, headers, widths)

    # Only records with a monthly premium (real or derived)
    monthly_records = [r for r in records if r.premium_monthly is not None]
    monthly_records.sort(key=lambda r: r.premium_monthly or 0, reverse=True)

    for i, r in enumerate(monthly_records, start=2):
        row = [
            r.insurance_type,
            r.insurer_company,
            r.policy_number,
            r.insured_name,
            r.premium_monthly,
            "Yes" if r.premium_monthly_derived else "No",
            r.currency,
            r.renewal_date,
            "Yes" if r.needs_review else "No",
        ]
        _write_row(ws, i, row, alt=(i % 2 == 0),
                   highlight=REVIEW_FILL if r.needs_review else None)
        ws.cell(row=i, column=5).number_format = '#,##0.00'

    # Summary row
    if monthly_records:
        total_row = len(monthly_records) + 2
        total_col = 5  # Monthly column
        ws.cell(row=total_row, column=1).value = "TOTAL"
        ws.cell(row=total_row, column=1).font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=total_col)
        total_cell.value = sum(r.premium_monthly or 0 for r in monthly_records)
        total_cell.number_format = '#,##0.00'
        total_cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Sheet: Needs Review
# ---------------------------------------------------------------------------

_REVIEW_HEADERS = [
    ("Type", 14),
    ("Insurer", 22),
    ("Policy #", 16),
    ("Insured Name", 22),
    ("Ext. Confidence", 14),
    ("Class. Confidence", 16),
    ("Missing Fields", 35),
    ("Notes", 40),
    ("Source Subject", 35),
    ("Record ID", 16),
]


def _build_needs_review(wb: Workbook, records: list[InsuranceRecord]) -> None:
    ws = wb.create_sheet("Needs Review")
    headers = [h for h, _ in _REVIEW_HEADERS]
    widths = [w for _, w in _REVIEW_HEADERS]
    _write_headers(ws, headers, widths)

    review_records = [r for r in records if r.needs_review]
    for i, r in enumerate(review_records, start=2):
        missing = _missing_fields(r)
        row = [
            r.insurance_type,
            r.insurer_company,
            r.policy_number,
            r.insured_name,
            r.extraction_confidence,
            r.classification_confidence,
            ", ".join(missing),
            r.extraction_notes,
            r.source_email_subject,
            r.record_id,
        ]
        _write_row(ws, i, row, alt=False, highlight=REVIEW_FILL)
        ws.cell(row=i, column=5).number_format = "0.00%"
        ws.cell(row=i, column=6).number_format = "0.00%"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Sheet: Sync Log
# ---------------------------------------------------------------------------

_SYNC_LOG_HEADERS = [
    ("Run ID", 20),
    ("Status", 12),
    ("Started", 20),
    ("Finished", 20),
    ("Messages Found", 14),
    ("New", 8),
    ("Skipped", 10),
    ("Attachments", 12),
    ("Records Created", 15),
    ("Needs Review", 13),
    ("Errors", 8),
]


def _build_sync_log(
    wb: Workbook,
    sync_runs: list[SyncRunSummary],
    errors: list[ProcessingError],
) -> None:
    ws = wb.create_sheet("Sync Log")
    headers = [h for h, _ in _SYNC_LOG_HEADERS]
    widths = [w for _, w in _SYNC_LOG_HEADERS]
    _write_headers(ws, headers, widths)

    for i, run in enumerate(reversed(sync_runs), start=2):
        row = [
            run.run_id,
            run.status,
            run.started_at,
            run.finished_at or "",
            run.messages_found,
            run.messages_new,
            run.messages_skipped,
            run.attachments_processed,
            run.records_created,
            run.records_needs_review,
            run.errors,
        ]
        highlight = ERROR_FILL if run.status == "failed" else (
            ALT_FILL if i % 2 == 0 else None
        )
        _write_row(ws, i, row, alt=False, highlight=highlight)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _write_headers(ws, headers: list[str], widths: list[int]) -> None:
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30


def _write_row(
    ws,
    row_idx: int,
    values: list[Any],
    alt: bool = False,
    highlight: Optional[PatternFill] = None,
) -> None:
    fill = highlight or (ALT_FILL if alt else None)
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col, value=value)
        cell.font = DATA_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=False)
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill


def _missing_fields(r: InsuranceRecord) -> list[str]:
    """Return list of important empty fields for this record."""
    checks = [
        ("insurer_company", "Insurer"),
        ("policy_number", "Policy #"),
        ("start_date", "Start Date"),
        ("end_date", "End Date"),
        ("premium_monthly", "Monthly Premium"),
    ]
    return [label for attr, label in checks if not getattr(r, attr)]
